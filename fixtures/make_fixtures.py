"""Build the two sample workbooks and matching JSON snapshots used by the tests.

    python fixtures/make_fixtures.py

Ventas_Sep.xlsx        one month per file (the "overwrite" / "copy last month's file" habits)
                       -> duplicate invoice numbers, amounts stored as text, a USD amount in an ARS column, a blank ID
Ventas_Acumulado.xlsx  rows appended month after month, Jan-Sep 2026
                       -> September looks ~35% down, because a third of its amounts were pasted as text

The JSON files mimic what Office.js returns (values / text / numberFormat) with Excel set to Spanish (Argentina),
so the TypeScript checks are tested against the same shapes they see inside Excel.
"""
import json
import random
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).parent
HEADERS = ["Nro Factura", "Fecha", "Cliente", "Producto", "Cantidad", "Importe"]
ARS_FMT = '"$"#,##0.00'
USD_FMT = '"US$"#,##0.00'
DATE_FMT = "dd/mm/yyyy"
CLIENTES = ["Ferretería Norte", "Distribuidora Sur", "Kiosco 24", "Almacén Don José", "Mayorista Oeste",
            "Panadería La Espiga", "Corralón Central", "Librería Belgrano"]
PRODUCTOS = ["Tornillos x100", "Pintura 4L", "Cable 2.5mm", "Cemento 50kg", "Lija fina", "Silicona", "Taladro"]


def ar_number(x: float, decimals: int = 2) -> str:
    s = f"{abs(x):,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-{s}" if x < 0 else s


def serial(d: date) -> float:
    return (datetime(d.year, d.month, d.day) - datetime(1899, 12, 30)).days


def display(v, fmt):
    if isinstance(v, str):
        return v
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if fmt == ARS_FMT:
        return "$" + ar_number(v)
    if fmt == USD_FMT:
        return "US$" + ar_number(v)
    if isinstance(v, int):
        return str(v)
    return ar_number(v)


def make_rows(rng, year, month, n, start_no, scale=1.0):
    rows = []
    for i in range(n):
        day = min(28, 1 + int(i * 28 / n)) if month != 9 else min(30, 1 + int(i * 30 / n))
        qty = rng.randint(8, 12)
        amount = round(qty * rng.uniform(2300, 2700) * scale, 2)
        rows.append([f"FA-0001-{start_no + i:08d}", date(year, month, day), rng.choice(CLIENTES),
                     rng.choice(PRODUCTOS), qty, amount])
    return rows


def write(path: Path, sheet: str, table_name: str, rows, formats, text_cells=()):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = formats.get((r - 2, c - 1))
            if fmt:
                cell.number_format = fmt
            elif isinstance(v, date):
                cell.number_format = DATE_FMT
            elif c == 6 and not isinstance(v, str):
                cell.number_format = ARS_FMT
    last = len(rows) + 1
    ref = f"A1:F{last}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for col, w in zip("ABCDEF", [20, 12, 22, 16, 10, 16]):
        ws.column_dimensions[col].width = w
    wb.save(path)

    # JSON snapshot, shaped like Office.js output.
    json_rows = []
    for r, row in enumerate(rows):
        out = []
        for c, v in enumerate(row):
            fmt = formats.get((r, c))
            if fmt is None:
                fmt = DATE_FMT if isinstance(v, date) else (ARS_FMT if c == 5 and not isinstance(v, str) else "General")
            if isinstance(v, str):
                fmt = formats.get((r, c), "General")
            val = serial(v) if isinstance(v, date) else (None if v is None else v)
            out.append({"v": val, "t": "" if v is None else display(v, fmt), "f": fmt})
        json_rows.append(out)
    snap = {"tableName": table_name, "sheet": sheet, "address": f"{sheet}!{ref}", "headers": HEADERS,
            "rows": json_rows, "firstRow": 1, "firstCol": 0}
    path.with_suffix(".json").write_text(json.dumps(snap, ensure_ascii=False, indent=0), encoding="utf-8")


def ventas_sep():
    rng = random.Random(7)
    rows = make_rows(rng, 2026, 9, 120, 1200)
    formats = {}
    # Duplicates: exact copy, trailing space, lower case.
    rows[40][0] = rows[12][0]
    rows[77][0] = rows[5][0] + " "
    rows[90][0] = rows[33][0].lower()
    # Blank ID.
    rows[101][0] = None
    # Amounts pasted as text, the way they arrive from e-mails and other systems.
    for r, txt in {8: "$ 12.345,50", 19: "12.345,50", 27: "8.900", 45: "$ 1.234,00", 63: "15.000,00",
                   88: "(2.500,00)", 110: "4.321,10"}.items():
        rows[r][5] = txt
    # A USD amount typed into the ARS column, plus one real number formatted as USD.
    rows[52][5] = "US$ 1,200.00"
    rows[70][5] = 950.0
    formats[(70, 5)] = USD_FMT
    write(HERE / "Ventas_Sep.xlsx", "Ventas", "Ventas", rows, formats)


def ventas_acumulado():
    rng = random.Random(11)
    rows, no = [], 5000
    for m in range(1, 10):
        n = 55 + m  # slow growth in volume
        month_rows = make_rows(rng, 2026, m, n, no, scale=1 + 0.02 * m)
        no += n
        if m == 9:
            for i in range(0, n, 3):  # every third amount pasted as text
                month_rows[i][5] = "$ " + ar_number(month_rows[i][5])
        rows.extend(month_rows)
    write(HERE / "Ventas_Acumulado.xlsx", "Acumulado", "VentasAcumulado", rows, {})


if __name__ == "__main__":
    ventas_sep()
    ventas_acumulado()
    print("wrote", *sorted(p.name for p in HERE.glob("Ventas_*")))
